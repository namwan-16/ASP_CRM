import csv
import io
import re
import zipfile
from datetime import date, datetime

from django.db import connection, transaction
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

from .models import Guardian, Student, StudentGuardian
from .serializers import GuardianSerializer, StudentSerializer, normalize_phone


HEADERS = [
    "student_id", "first_name", "last_name", "date_of_birth", "school", "year_level", "subject",
    "email", "phone", "permission_to_travel_alone", "is_active", "medical_information", "special_circumstances",
    "guardian_first_name", "guardian_last_name", "guardian_phone", "guardian_email", "relationship",
    "is_primary_contact", "is_emergency_contact",
]
REQUIRED = {"first_name", "last_name", "date_of_birth", "guardian_first_name", "guardian_last_name", "guardian_phone"}
MAX_ROWS = 5000
MAX_BYTES = 10 * 1024 * 1024


def read_rows(upload):
    if upload.size > MAX_BYTES:
        raise ValidationError({"file": "File size must be 10 MB or less."})
    extension = upload.name.rsplit(".", 1)[-1].lower()
    workbook = None
    try:
        if extension == "csv":
            contents = upload.read(MAX_BYTES + 1).decode("utf-8-sig")
            table = csv.reader(io.StringIO(contents), strict=True)
        elif extension == "xlsx":
            with zipfile.ZipFile(upload) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValidationError({"file": "Workbook expands beyond the 50 MB safety limit."})
            upload.seek(0)
            workbook = load_workbook(upload, read_only=True, data_only=False, keep_links=False)
            sheet = workbook.active
            # Ignore potentially incorrect worksheet dimension metadata.
            sheet.reset_dimensions()

            def excel_rows():
                for row in sheet.iter_rows():
                    if len(row) > 40:
                        raise ValidationError({"file": "A workbook may contain at most 40 columns."})
                    if any(cell.data_type == "f" for cell in row):
                        raise ValidationError({"file": "Formulas are not allowed. Paste values before importing."})
                    yield [cell.value for cell in row]
            table = excel_rows()
        else:
            raise ValidationError({"file": "Upload an .xlsx or UTF-8 .csv file. Legacy .xls files must be saved as .xlsx."})

        header_row = next(table, None)
        if not header_row or len(header_row) > 40:
            raise ValidationError({"file": "The file needs a header row with at most 40 columns."})
        headers = [re.sub(r"[\s-]+", "_", str(cell or "").strip().lower()) for cell in header_row]
        if len(headers) != len(set(headers)):
            raise ValidationError({"file": "Duplicate column headers are not allowed."})
        if REQUIRED - set(headers):
            raise ValidationError({"file": "Missing columns: " + ", ".join(sorted(REQUIRED - set(headers)))})
        if set(headers) - set(HEADERS):
            raise ValidationError({"file": "Unknown columns: " + ", ".join(sorted(set(headers) - set(HEADERS)))})
        rows = []
        for number, values in enumerate(table, 2):
            if number > MAX_ROWS + 1:
                raise ValidationError({"file": f"Import at most {MAX_ROWS} rows at a time."})
            if not any(value not in (None, "") for value in values):
                continue
            if len(values) > len(headers) and any(value not in (None, "") for value in values[len(headers):]):
                raise ValidationError({"file": f"Row {number} has more values than headers."})
            values = list(values) + [None] * max(0, len(headers) - len(values))
            rows.append((number, dict(zip(headers, values))))
        if not rows:
            raise ValidationError({"file": "The file contains no data rows."})
        return rows
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError({"file": "Unable to read this file. Check its format, encoding and contents."}) from exc
    finally:
        if workbook:
            workbook.close()


def text(value):
    return "" if value is None else str(value).strip()


def birth_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    for pattern in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text(value), pattern).date().isoformat()
        except ValueError:
            pass
    raise ValidationError({"date_of_birth": "Use YYYY-MM-DD or DD/MM/YYYY."})


def boolean(value, default):
    if value in (None, ""):
        return default
    value = text(value).lower()
    if value in ("true", "yes", "1", "1.0"):
        return True
    if value in ("false", "no", "0", "0.0"):
        return False
    raise ValidationError("Boolean columns accept Yes/No, True/False or 1/0.")


def import_row(row, seen, guardian_seen):
    for field in ("guardian_phone", "phone"):
        if row.get(field) not in (None, "") and not isinstance(row[field], str):
            raise ValidationError({field: "Store phone numbers as text in Excel to preserve leading zeroes."})
    phone = normalize_phone(text(row["guardian_phone"]))
    guardian_data = {"first_name": text(row["guardian_first_name"]), "last_name": text(row["guardian_last_name"]), "phone": phone}
    if text(row.get("guardian_email")):
        guardian_data["email"] = text(row["guardian_email"])
    guardian_identity = (guardian_data["first_name"].casefold(), guardian_data["last_name"].casefold())
    signature = (guardian_identity, guardian_data.get("email", "").casefold())
    if phone in guardian_seen and guardian_seen[phone] != signature:
        raise ValidationError({"guardian_phone": "Conflicting guardian details for the same phone number in this file."})
    guardian_seen[phone] = signature
    matches = list(Guardian.objects.select_for_update().filter(phone=phone))
    if len(matches) > 1:
        raise ValidationError({"guardian_phone": "Multiple existing guardians share this phone. Resolve them before importing."})
    guardian = matches[0] if matches else None
    if guardian and (guardian.first_name.casefold(), guardian.last_name.casefold()) != guardian_identity:
        raise ValidationError({"guardian_phone": "This phone belongs to a different guardian. Review the row."})
    guardian_serializer = GuardianSerializer(guardian, data=guardian_data, partial=bool(guardian))
    guardian_serializer.is_valid(raise_exception=True)

    student_data = {key: text(row[key]) for key in (
        "first_name", "last_name", "school", "year_level", "subject", "email", "phone", "medical_information", "special_circumstances"
    ) if key in row and (key in ("first_name", "last_name") or text(row[key]))}
    student_data["date_of_birth"] = birth_date(row["date_of_birth"])
    for key, default in (("permission_to_travel_alone", False), ("is_active", True)):
        if row.get(key) not in (None, ""):
            student_data[key] = boolean(row[key], default)
    student = None
    if text(row.get("student_id")):
        try:
            student_id = int(text(row["student_id"]))
        except ValueError as exc:
            raise ValidationError({"student_id": "Use the numeric student ID from the CRM."}) from exc
        student = Student.objects.select_for_update().filter(pk=student_id).first()
        if not student:
            raise ValidationError({"student_id": "This student ID does not exist."})
    elif guardian:
        matches = list(Student.objects.select_for_update().filter(
            guardian_links__guardian=guardian, first_name__iexact=student_data["first_name"], last_name__iexact=student_data["last_name"],
        ))
        exact = [match for match in matches if match.date_of_birth.isoformat() == student_data["date_of_birth"]]
        if len(exact) > 1 or (matches and not exact):
            raise ValidationError({"student_id": "Ambiguous student match or changed birth date. Supply the existing student ID."})
        student = exact[0] if exact else None
    identity = ("id", student.pk) if student else ("new", phone, student_data["first_name"].casefold(), student_data["last_name"].casefold(), student_data["date_of_birth"])
    # Detect repeated identities even when an earlier row created the student.
    natural = (phone, student_data["first_name"].casefold(), student_data["last_name"].casefold(), student_data["date_of_birth"])
    if identity in seen or natural in seen:
        raise ValidationError("Duplicate student row in this file.")
    seen.update((identity, natural))
    serializer = StudentSerializer(student, data=student_data, partial=bool(student))
    serializer.is_valid(raise_exception=True)
    created = student is None
    guardian = guardian_serializer.save()
    student = serializer.save()
    primary = boolean(row.get("is_primary_contact"), not StudentGuardian.objects.filter(student=student, is_primary_contact=True).exclude(guardian=guardian).exists())
    if primary and StudentGuardian.objects.filter(student=student, is_primary_contact=True).exclude(guardian=guardian).exists():
        raise ValidationError({"is_primary_contact": "This student already has another primary guardian."})
    relationship = text(row.get("relationship")) or "Guardian"
    if len(relationship) > 50:
        raise ValidationError({"relationship": "Use at most 50 characters."})
    StudentGuardian.objects.update_or_create(student=student, guardian=guardian, defaults={
        "relationship": relationship, "is_primary_contact": primary,
        "is_emergency_contact": boolean(row.get("is_emergency_contact"), True),
    })
    return created


def import_students(upload, dry_run=False):
    rows = read_rows(upload)
    errors, seen, guardian_seen = [], set(), {}
    created = updated = 0
    with transaction.atomic():
        if connection.vendor == "postgresql":
            with connection.cursor() as cursor:
                cursor.execute("SELECT pg_advisory_xact_lock(%s)", [6202026])
        for number, row in rows:
            try:
                with transaction.atomic():
                    was_created = import_row(row, seen, guardian_seen)
                created += int(was_created)
                updated += int(not was_created)
            except ValidationError as exc:
                errors.append({"row": number, "errors": exc.detail})
        if errors:
            raise ValidationError({"detail": "No records saved. Fix the rows below and try again.", "errors": errors})
        if dry_run:
            transaction.set_rollback(True)
    return {"dry_run": dry_run, "rows": len(rows), "created": created, "updated": updated}
