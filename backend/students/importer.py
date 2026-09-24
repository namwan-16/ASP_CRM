import re
from datetime import date, datetime

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook

from .models import Student


HEADER_ALIASES = {
    "phonenumber": "phone_number",
    "phone": "phone_number",
    "mobile": "phone_number",
    "mobilenumber": "phone_number",
    "firstname": "first_name",
    "givenname": "first_name",
    "lastname": "last_name",
    "surname": "last_name",
    "email": "email",
    "dateofbirth": "date_of_birth",
    "dob": "date_of_birth",
    "school": "school",
    "yearlevel": "year_level",
    "grade": "year_level",
    "guardianname": "guardian_name",
    "guardianphone": "guardian_phone",
    "address": "address",
    "notes": "notes",
    "active": "is_active",
}

REQUIRED_FIELDS = {
    "phone_number",
    "first_name",
    "last_name",
}


def normalise_header(value):
    key = re.sub(
        r"[^a-z0-9]",
        "",
        str(value or "").strip().lower(),
    )

    return HEADER_ALIASES.get(key)


def cell_text(value):
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def normalise_phone(value, required=True):
    raw = cell_text(value)

    if not raw:
        if required:
            raise ValueError("Phone Number is required.")

        return ""

    digits = re.sub(r"\D", "", raw)

    # Change 0061 to 61
    if digits.startswith("0061"):
        digits = digits[2:]

    # Convert +61 Australian number into local 0 format.
    # Example: +61 412 345 678 becomes 0412345678.
    if digits.startswith("61") and len(digits) == 11:
        digits = "0" + digits[2:]

    # Restore the leading zero if Excel removed it.
    elif len(digits) == 9 and digits.startswith("4"):
        digits = "0" + digits

    if not 8 <= len(digits) <= 15:
        raise ValueError(
            "Phone Number must contain between 8 and 15 digits."
        )

    if raw.startswith("+") and not digits.startswith("0"):
        return f"+{digits}"

    return digits


def parse_date(value):
    if value in (None, ""):
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = cell_text(value)

    formats = (
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d-%m-%Y",
    )

    for date_format in formats:
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue

    raise ValueError(
        "Date of Birth must use YYYY-MM-DD or DD/MM/YYYY."
    )


def parse_boolean(value):
    if value in (None, ""):
        return True

    if isinstance(value, bool):
        return value

    text = cell_text(value).lower()

    if text in {"yes", "true", "1", "active"}:
        return True

    if text in {"no", "false", "0", "inactive"}:
        return False

    raise ValueError(
        "Active must be Yes, No, True, False, 1, or 0."
    )


def format_validation_error(error):
    if hasattr(error, "message_dict"):
        return "; ".join(
            f"{field}: {', '.join(messages)}"
            for field, messages in error.message_dict.items()
        )

    return "; ".join(error.messages)


def import_students(uploaded_file):
    workbook = load_workbook(
        uploaded_file,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)

    try:
        raw_headers = next(rows)
    except StopIteration:
        raise ValueError("The spreadsheet is empty.")

    column_map = {}

    for column_index, raw_header in enumerate(raw_headers):
        field_name = normalise_header(raw_header)

        if field_name:
            column_map[column_index] = field_name

    missing_fields = REQUIRED_FIELDS - set(column_map.values())

    if missing_fields:
        names = {
            "phone_number": "Phone Number",
            "first_name": "First Name",
            "last_name": "Last Name",
        }

        missing_text = ", ".join(
            names[field] for field in sorted(missing_fields)
        )

        raise ValueError(
            f"Missing required column(s): {missing_text}."
        )

    result = {
        "created": 0,
        "updated": 0,
        "failed": 0,
        "skipped_blank": 0,
        "errors": [],
    }

    for row_number, row in enumerate(rows, start=2):
        values = {
            field_name: (
                row[column_index]
                if column_index < len(row)
                else None
            )
            for column_index, field_name in column_map.items()
        }

        if not any(
            value not in (None, "")
            for value in values.values()
        ):
            result["skipped_blank"] += 1
            continue

        try:
            phone_number = normalise_phone(
                values.get("phone_number")
            )

            first_name = cell_text(
                values.get("first_name")
            )

            last_name = cell_text(
                values.get("last_name")
            )

            if not first_name:
                raise ValueError("First Name is required.")

            if not last_name:
                raise ValueError("Last Name is required.")

            student_data = {
                "first_name": first_name,
                "last_name": last_name,
                "email": cell_text(values.get("email")),
                "date_of_birth": parse_date(
                    values.get("date_of_birth")
                ),
                "school": cell_text(values.get("school")),
                "year_level": cell_text(
                    values.get("year_level")
                ),
                "guardian_name": cell_text(
                    values.get("guardian_name")
                ),
                "guardian_phone": normalise_phone(
                    values.get("guardian_phone"),
                    required=False,
                ),
                "address": cell_text(values.get("address")),
                "notes": cell_text(values.get("notes")),
                "is_active": parse_boolean(
                    values.get("is_active")
                ),
            }

            with transaction.atomic():
                student = Student.objects.filter(
                    phone_number=phone_number
                ).first()

                student_created = student is None

                if student_created:
                    student = Student(
                        phone_number=phone_number
                    )

                for field, value in student_data.items():
                    setattr(student, field, value)

                student.full_clean()
                student.save()

            if student_created:
                result["created"] += 1
            else:
                result["updated"] += 1

        except (ValueError, ValidationError) as error:
            result["failed"] += 1

            if isinstance(error, ValidationError):
                message = format_validation_error(error)
            else:
                message = str(error)

            result["errors"].append(
                {
                    "row": row_number,
                    "message": message,
                }
            )

    return result